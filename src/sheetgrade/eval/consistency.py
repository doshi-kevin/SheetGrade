"""Independent checks that a mutation's manifest matches what actually happened.

Part 3 gave every mutation a receipt (`MutationRecord`) claiming what it did.
Nothing has ever checked whether that receipt is telling the truth against the
actual before/after workbooks. A checker here re-derives the change from the
IR directly and deliberately avoids reusing any helper the mutation itself
used to make its claim (see `wrong_input_correct_method`'s reference-range
expansion, duplicated here rather than imported) -- if the checker shared
logic with the thing it's checking, a bug in that shared logic would be
invisible to both.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass

from openpyxl.utils import get_column_letter, range_boundaries

from sheetgrade.ir import CellAddress, CellType, Sheet, Workbook
from sheetgrade.mutate.manifest import MutationRecord
from sheetgrade.reader import a1_to_address, address_to_a1


@dataclass(frozen=True, slots=True)
class ConsistencyResult:
    mutation: str
    sheet: str
    ok: bool
    reason: str


def _ok(record: MutationRecord) -> ConsistencyResult:
    return ConsistencyResult(record.mutation, record.sheet, True, "")


def _fail(record: MutationRecord, reason: str) -> ConsistencyResult:
    return ConsistencyResult(record.mutation, record.sheet, False, reason)


def _sheet(workbook: Workbook, name: str) -> Sheet:
    for sheet in workbook.sheets:
        if sheet.name == name:
            return sheet
    raise KeyError(name)


def _moved(before: Workbook, after: Workbook, record: MutationRecord, move: Callable[[CellAddress], CellAddress]) -> ConsistencyResult:
    """Shared shape for every mutation that only relocates cells: every cell in
    `before` must reappear at `move(address)` in `after`, value untouched."""
    b, a = _sheet(before, record.sheet), _sheet(after, record.sheet)
    for addr, cell in b.cells.items():
        new_addr = move(addr)
        moved_cell = a.cells.get(new_addr)
        if moved_cell is None or moved_cell.value != cell.value:
            return _fail(record, f"{address_to_a1(addr)}={cell.value!r} did not land at {address_to_a1(new_addr)}")
    return _ok(record)


def _check_insert_column(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    at_col = record.details["at_col"]
    a = _sheet(after, record.sheet)
    b = _sheet(before, record.sheet)
    if a.n_cols != b.n_cols + 1:
        return _fail(record, f"n_cols went {b.n_cols} -> {a.n_cols}, expected +1")
    return _moved(before, after, record, lambda addr: CellAddress(addr.row, addr.col + 1 if addr.col >= at_col else addr.col))


def _check_delete_row(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    """delete_row's `deleted_cells` list is the harness's ground truth for "no
    counterpart exists" -- later alignment metrics (Part 11/12) will check that
    the aligner correctly abstains on exactly these cells rather than forcing
    a match."""
    at_row = record.details["at_row"]
    claimed_deleted = set(record.details["deleted_cells"])
    b, a = _sheet(before, record.sheet), _sheet(after, record.sheet)
    if a.n_rows != b.n_rows - 1:
        return _fail(record, f"n_rows went {b.n_rows} -> {a.n_rows}, expected -1")
    actual_deleted = {address_to_a1(addr) for addr in b.cells if addr.row == at_row}
    if actual_deleted != claimed_deleted:
        return _fail(record, f"manifest claims {claimed_deleted} deleted, row {at_row} actually held {actual_deleted}")

    for addr, cell in b.cells.items():
        if addr.row == at_row:
            continue
        new_addr = CellAddress(addr.row - 1 if addr.row > at_row else addr.row, addr.col)
        moved_cell = a.cells.get(new_addr)
        if moved_cell is None or moved_cell.value != cell.value:
            return _fail(record, f"{address_to_a1(addr)}={cell.value!r} did not land at {address_to_a1(new_addr)}")
    return _ok(record)


def _check_rename_header(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    addr = a1_to_address(record.details["address"])
    b, a = _sheet(before, record.sheet), _sheet(after, record.sheet)
    if b.cells[addr].value != record.details["before"]:
        return _fail(record, f"manifest 'before' ({record.details['before']!r}) doesn't match actual before value ({b.cells[addr].value!r})")
    if a.cells[addr].value != record.details["after"]:
        return _fail(record, f"manifest 'after' ({record.details['after']!r}) doesn't match actual after value ({a.cells[addr].value!r})")
    for other_addr, cell in b.cells.items():
        if other_addr == addr:
            continue
        other_after = a.cells.get(other_addr)
        if other_after is None or other_after.value != cell.value:
            return _fail(record, f"an unrelated cell {address_to_a1(other_addr)} changed too")
    return _ok(record)


def _check_reorder_columns(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    mapping = record.details["old_to_new"]
    b, a = _sheet(before, record.sheet), _sheet(after, record.sheet)
    if len(a.cells) != len(b.cells):
        return _fail(record, f"cell count changed: {len(b.cells)} -> {len(a.cells)}")
    return _moved(before, after, record, lambda addr: CellAddress(addr.row, mapping[addr.col]))


def _check_shift_region(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    min_row, min_col, max_row, max_col = record.details["box"]
    drow, dcol = record.details["offset"]

    def move(addr: CellAddress) -> CellAddress:
        if min_row <= addr.row <= max_row and min_col <= addr.col <= max_col:
            return CellAddress(addr.row + drow, addr.col + dcol)
        return addr

    return _moved(before, after, record, move)


def _check_recolour(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    addr = a1_to_address(record.details["address"])
    b, a = _sheet(before, record.sheet), _sheet(after, record.sheet)
    if b.cells[addr].fill_color_rgb != record.details["before"]:
        return _fail(record, "manifest 'before' fill colour doesn't match the actual before fill colour")
    if a.cells[addr].fill_color_rgb != record.details["after"]:
        return _fail(record, "manifest 'after' fill colour doesn't match the actual after fill colour")
    if a.cells[addr].value != b.cells[addr].value:
        return _fail(record, "recolour changed the cell's value, not just its colour")
    return _ok(record)


def _check_hardcode_formula(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    addr = a1_to_address(record.details["address"])
    b, a = _sheet(before, record.sheet), _sheet(after, record.sheet)
    if b.cells[addr].formula != record.details["old_formula"]:
        return _fail(record, "manifest 'old_formula' doesn't match the actual before formula")
    if a.cells[addr].formula is not None:
        return _fail(record, "formula is still present after hardcode_formula")
    if a.cells[addr].value != record.details["hardcoded_value"]:
        return _fail(record, "manifest 'hardcoded_value' doesn't match the actual after value")
    return _ok(record)


def _check_perturb_value(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    addr = a1_to_address(record.details["address"])
    b, a = _sheet(before, record.sheet), _sheet(after, record.sheet)
    if b.cells[addr].value != record.details["before"]:
        return _fail(record, "manifest 'before' doesn't match the actual before value")
    if a.cells[addr].value != record.details["after"]:
        return _fail(record, "manifest 'after' doesn't match the actual after value")
    if a.cells[addr].data_type is not CellType.NUMBER:
        return _fail(record, "cell is no longer numeric after perturb_value")
    return _ok(record)


def _check_merge_cells(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    cell_range = record.details["range"]
    a = _sheet(after, record.sheet)
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    assert min_col and min_row and max_col and max_row
    top_left = CellAddress(min_row - 1, min_col - 1)
    anchor = a.cells.get(top_left)
    if anchor is None or anchor.merge_range != cell_range:
        return _fail(record, f"anchor cell {address_to_a1(top_left)} is missing or not marked with merge_range={cell_range}")
    for row in range(min_row - 1, max_row):
        for col in range(min_col - 1, max_col):
            addr = CellAddress(row, col)
            if addr != top_left and addr in a.cells:
                return _fail(record, f"{address_to_a1(addr)} should have been absorbed into the merge but is still a separate cell")
    return _ok(record)


def _check_rename_sheet(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    old_name, new_name = record.details["old_name"], record.details["new_name"]
    if any(s.name == old_name for s in after.sheets):
        return _fail(record, f"old sheet name {old_name!r} is still present after rename_sheet")
    if not any(s.name == new_name for s in after.sheets):
        return _fail(record, f"new sheet name {new_name!r} is missing after rename_sheet")
    old_sheet = _sheet(before, old_name)
    new_sheet = _sheet(after, new_name)
    if old_sheet.cells != new_sheet.cells:
        return _fail(record, "cell contents changed during what should have been a pure rename")
    return _ok(record)


def _check_split_table(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    at_row, gap_rows = record.details["at_row"], record.details["gap_rows"]
    b, a = _sheet(before, record.sheet), _sheet(after, record.sheet)
    if a.n_rows != b.n_rows + gap_rows:
        return _fail(record, f"n_rows went {b.n_rows} -> {a.n_rows}, expected +{gap_rows}")
    return _moved(before, after, record, lambda addr: CellAddress(addr.row + gap_rows if addr.row >= at_row else addr.row, addr.col))


def _check_add_scratch_work(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    addr = a1_to_address(record.details["address"])
    b, a = _sheet(before, record.sheet), _sheet(after, record.sheet)
    new_cell = a.cells.get(addr)
    if new_cell is None or new_cell.value != record.details["text"]:
        return _fail(record, "scratch-work cell is missing or doesn't match the manifest text")
    for other_addr, cell in b.cells.items():
        other_after = a.cells.get(other_addr)
        if other_after is None or other_after.value != cell.value:
            return _fail(record, f"an existing cell {address_to_a1(other_addr)} was disturbed by add_scratch_work")
    return _ok(record)


def _expand_ref(ref: str) -> set[str]:
    """Independent copy of range expansion: 'B3:E3' -> {B3, C3, D3, E3}."""
    ref = ref.split("!")[-1]
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    assert min_col and min_row and max_col and max_row
    return {
        f"{get_column_letter(c)}{r}"
        for r in range(min_row, max_row + 1)
        for c in range(min_col, max_col + 1)
    }


def _check_wrong_input_correct_method(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    input_a1 = record.details["input_address"]
    addr = a1_to_address(input_a1)
    b, a = _sheet(before, record.sheet), _sheet(after, record.sheet)
    if b.cells[addr].value != record.details["before"]:
        return _fail(record, "manifest 'before' doesn't match the actual before value")
    if a.cells[addr].value != record.details["after"]:
        return _fail(record, "manifest 'after' doesn't match the actual after value")

    actual_dependents = {
        key
        for key, refs in before.dependency_graph.items()
        if key.startswith(f"{record.sheet}!")
        for ref in refs
        if input_a1 in _expand_ref(ref)
    }
    claimed = set(record.details["formulas_correct_given_bad_input"])
    if claimed != actual_dependents:
        return _fail(record, f"manifest claims dependents {claimed}, dependency graph says {actual_dependents}")

    for key in claimed:
        formula_addr = a1_to_address(key.split("!")[1])
        if a.cells[formula_addr].formula != b.cells[formula_addr].formula:
            return _fail(record, f"{key}'s formula text changed -- it's supposed to still be correct, just fed a bad input")
    return _ok(record)


_CHECKS: dict[str, Callable[[Workbook, Workbook, MutationRecord], ConsistencyResult]] = {
    "insert_column": _check_insert_column,
    "delete_row": _check_delete_row,
    "rename_header": _check_rename_header,
    "reorder_columns": _check_reorder_columns,
    "shift_region": _check_shift_region,
    "recolour": _check_recolour,
    "hardcode_formula": _check_hardcode_formula,
    "perturb_value": _check_perturb_value,
    "merge_cells": _check_merge_cells,
    "rename_sheet": _check_rename_sheet,
    "split_table": _check_split_table,
    "add_scratch_work": _check_add_scratch_work,
    "wrong_input_correct_method": _check_wrong_input_correct_method,
}


def check_step(before: Workbook, after: Workbook, record: MutationRecord) -> ConsistencyResult:
    checker = _CHECKS.get(record.mutation)
    if checker is None:
        return _fail(record, f"no consistency checker registered for mutation {record.mutation!r}")
    return checker(before, after, record)

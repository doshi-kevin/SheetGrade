"""The 13 mutation types from PROJECT_PLAN.md Part 3.

Every function takes (workbook, rng, ...explicit params...) and returns a
brand-new Workbook plus a MutationRecord describing what changed. `rng` is a
seeded random.Random -- functions that weren't given an explicit target pick
one using `rng`, so the same seed always produces the same mutation.

Known, deliberate limitation: structural mutations (insert_column, delete_row,
reorder_columns, shift_region, split_table) move *cell positions* but do not
rewrite formula *text*. A formula like `=B3-B4` still says B3 and B4 after a
column is inserted before B, even though the values that were in B3/B4 have
physically moved. Fixing that requires understanding formula text well enough
to rewrite it -- which needs the AST/parser work deferred to Part 14. Logged
in DECISIONS.md rather than silently ignored.
"""

from __future__ import annotations

import random
from collections.abc import Callable
from dataclasses import replace

from openpyxl.utils import get_column_letter, range_boundaries

from sheetgrade.ir import Cell, CellAddress, CellType, Workbook
from sheetgrade.mutate.helpers import get_sheet, replace_sheet, set_cell
from sheetgrade.mutate.manifest import MutationRecord
from sheetgrade.reader import address_to_a1, build_dependency_graph

_RECOLOUR_PALETTE = ("FFFF0000", "FF00B050", "FF0070C0", "FFFFFF00")
_SCRATCH_TEXTS = ("scratch calc", "workspace - ignore", "notes to self", "temp")


def insert_column(workbook: Workbook, rng: random.Random, sheet: str, at_col: int | None = None) -> tuple[Workbook, MutationRecord]:
    sh = get_sheet(workbook, sheet)
    if at_col is None:
        at_col = rng.randint(0, sh.n_cols)

    new_cells: dict[CellAddress, Cell] = {}
    for addr, cell in sh.cells.items():
        if addr.col >= at_col:
            new_addr = CellAddress(addr.row, addr.col + 1)
            new_cells[new_addr] = replace(cell, address=new_addr)
        else:
            new_cells[addr] = cell

    new_sheet = replace(sh, cells=new_cells, n_cols=sh.n_cols + 1)
    record = MutationRecord("insert_column", sheet, "structural_only", {"at_col": at_col})
    return replace_sheet(workbook, sheet, new_sheet), record


def delete_row(workbook: Workbook, rng: random.Random, sheet: str, at_row: int | None = None) -> tuple[Workbook, MutationRecord]:
    sh = get_sheet(workbook, sheet)
    if at_row is None:
        at_row = rng.randint(0, sh.n_rows - 1)

    deleted = [address_to_a1(a) for a in sh.cells if a.row == at_row]
    new_cells: dict[CellAddress, Cell] = {}
    for addr, cell in sh.cells.items():
        if addr.row == at_row:
            continue
        if addr.row > at_row:
            new_addr = CellAddress(addr.row - 1, addr.col)
            new_cells[new_addr] = replace(cell, address=new_addr)
        else:
            new_cells[addr] = cell

    new_sheet = replace(sh, cells=new_cells, n_rows=sh.n_rows - 1)
    record = MutationRecord(
        "delete_row", sheet, "structural_only", {"at_row": at_row, "deleted_cells": deleted}
    )
    return replace_sheet(workbook, sheet, new_sheet), record


def rename_header(workbook: Workbook, rng: random.Random, sheet: str, address: CellAddress, new_name: str) -> tuple[Workbook, MutationRecord]:
    sh = get_sheet(workbook, sheet)
    cell = sh.cells[address]
    new_sheet = set_cell(sh, address, replace(cell, value=new_name))
    record = MutationRecord(
        "rename_header",
        sheet,
        "structural_only",
        {"address": address_to_a1(address), "before": cell.value, "after": new_name},
    )
    return replace_sheet(workbook, sheet, new_sheet), record


def reorder_columns(
    workbook: Workbook, rng: random.Random, sheet: str, old_to_new: list[int] | None = None
) -> tuple[Workbook, MutationRecord]:
    sh = get_sheet(workbook, sheet)
    if old_to_new is None:
        old_to_new = rng.sample(range(sh.n_cols), sh.n_cols)

    new_cells: dict[CellAddress, Cell] = {}
    for addr, cell in sh.cells.items():
        new_addr = CellAddress(addr.row, old_to_new[addr.col])
        new_cells[new_addr] = replace(cell, address=new_addr)

    new_sheet = replace(sh, cells=new_cells)
    record = MutationRecord("reorder_columns", sheet, "structural_only", {"old_to_new": old_to_new})
    return replace_sheet(workbook, sheet, new_sheet), record


def shift_region(
    workbook: Workbook,
    rng: random.Random,
    sheet: str,
    drow: int,
    dcol: int,
    min_row: int = 0,
    min_col: int = 0,
    max_row: int | None = None,
    max_col: int | None = None,
) -> tuple[Workbook, MutationRecord]:
    sh = get_sheet(workbook, sheet)
    max_row = sh.n_rows - 1 if max_row is None else max_row
    max_col = sh.n_cols - 1 if max_col is None else max_col

    new_cells: dict[CellAddress, Cell] = {}
    for addr, cell in sh.cells.items():
        if min_row <= addr.row <= max_row and min_col <= addr.col <= max_col:
            new_addr = CellAddress(addr.row + drow, addr.col + dcol)
            assert new_addr.row >= 0 and new_addr.col >= 0, "shift_region moved a cell off-sheet"
            new_cells[new_addr] = replace(cell, address=new_addr)
        else:
            new_cells[addr] = cell

    n_rows = max(sh.n_rows, max(a.row for a in new_cells) + 1) if new_cells else sh.n_rows
    n_cols = max(sh.n_cols, max(a.col for a in new_cells) + 1) if new_cells else sh.n_cols
    new_sheet = replace(sh, cells=new_cells, n_rows=n_rows, n_cols=n_cols)
    record = MutationRecord(
        "shift_region",
        sheet,
        "structural_only",
        {"box": [min_row, min_col, max_row, max_col], "offset": [drow, dcol]},
    )
    return replace_sheet(workbook, sheet, new_sheet), record


def recolour(
    workbook: Workbook,
    rng: random.Random,
    sheet: str,
    address: CellAddress | None = None,
    new_rgb: str | None = None,
) -> tuple[Workbook, MutationRecord]:
    sh = get_sheet(workbook, sheet)
    if address is None:
        address = rng.choice(list(sh.cells))
    if new_rgb is None:
        new_rgb = rng.choice(_RECOLOUR_PALETTE)

    cell = sh.cells[address]
    new_sheet = set_cell(sh, address, replace(cell, fill_color_rgb=new_rgb))
    record = MutationRecord(
        "recolour",
        sheet,
        "cosmetic_only",
        {"address": address_to_a1(address), "before": cell.fill_color_rgb, "after": new_rgb},
    )
    return replace_sheet(workbook, sheet, new_sheet), record


def hardcode_formula(workbook: Workbook, rng: random.Random, sheet: str, address: CellAddress) -> tuple[Workbook, MutationRecord]:
    sh = get_sheet(workbook, sheet)
    cell = sh.cells[address]
    if cell.data_type is not CellType.FORMULA:
        raise ValueError(f"{address} is not a formula cell")

    value = cell.computed_value if cell.computed_value is not None else 0
    new_type = CellType.NUMBER if isinstance(value, int | float) else CellType.STRING
    new_cell = replace(cell, data_type=new_type, value=value, formula=None, computed_value=None)
    new_sheet = set_cell(sh, address, new_cell)
    record = MutationRecord(
        "hardcode_formula",
        sheet,
        "method_lost",
        {"address": address_to_a1(address), "old_formula": cell.formula, "hardcoded_value": value},
    )
    return replace_sheet(workbook, sheet, new_sheet), record


def perturb_value(
    workbook: Workbook,
    rng: random.Random,
    sheet: str,
    address: CellAddress | None = None,
    delta: float | None = None,
) -> tuple[Workbook, MutationRecord]:
    sh = get_sheet(workbook, sheet)
    if address is None:
        candidates = [a for a, c in sh.cells.items() if c.data_type is CellType.NUMBER]
        address = rng.choice(candidates)
    if delta is None:
        delta = rng.choice([-1, 1]) * rng.randint(1, 10)

    cell = sh.cells[address]
    if cell.data_type is not CellType.NUMBER or not isinstance(cell.value, int | float):
        raise ValueError(f"{address} is not a numeric cell")
    new_value = cell.value + delta
    new_sheet = set_cell(sh, address, replace(cell, value=new_value))
    record = MutationRecord(
        "perturb_value",
        sheet,
        "value_changed",
        {"address": address_to_a1(address), "before": cell.value, "after": new_value},
    )
    return replace_sheet(workbook, sheet, new_sheet), record


def merge_cells(workbook: Workbook, rng: random.Random, sheet: str, cell_range: str) -> tuple[Workbook, MutationRecord]:
    sh = get_sheet(workbook, sheet)
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    assert min_col and min_row and max_col and max_row
    top_left = CellAddress(min_row - 1, min_col - 1)

    new_cells = dict(sh.cells)
    anchor = new_cells.get(top_left, Cell(address=top_left, value=None, data_type=CellType.BLANK))
    new_cells[top_left] = replace(anchor, merge_range=cell_range)
    for row in range(min_row - 1, max_row):
        for col in range(min_col - 1, max_col):
            addr = CellAddress(row, col)
            if addr != top_left:
                new_cells.pop(addr, None)

    new_sheet = replace(sh, cells=new_cells)
    record = MutationRecord("merge_cells", sheet, "structural_only", {"range": cell_range})
    return replace_sheet(workbook, sheet, new_sheet), record


def rename_sheet(workbook: Workbook, rng: random.Random, sheet: str, new_name: str) -> tuple[Workbook, MutationRecord]:
    new_sheets = tuple(replace(s, name=new_name) if s.name == sheet else s for s in workbook.sheets)
    new_workbook = replace(
        workbook, sheets=new_sheets, dependency_graph=build_dependency_graph(new_sheets)
    )
    record = MutationRecord("rename_sheet", sheet, "structural_only", {"old_name": sheet, "new_name": new_name})
    return new_workbook, record


def split_table(
    workbook: Workbook, rng: random.Random, sheet: str, at_row: int | None = None, gap_rows: int = 3
) -> tuple[Workbook, MutationRecord]:
    sh = get_sheet(workbook, sheet)
    if at_row is None:
        at_row = rng.randint(1, sh.n_rows - 1)

    new_cells: dict[CellAddress, Cell] = {}
    for addr, cell in sh.cells.items():
        if addr.row >= at_row:
            new_addr = CellAddress(addr.row + gap_rows, addr.col)
            new_cells[new_addr] = replace(cell, address=new_addr)
        else:
            new_cells[addr] = cell

    new_sheet = replace(sh, cells=new_cells, n_rows=sh.n_rows + gap_rows)
    record = MutationRecord(
        "split_table", sheet, "structural_only", {"at_row": at_row, "gap_rows": gap_rows}
    )
    return replace_sheet(workbook, sheet, new_sheet), record


def add_scratch_work(
    workbook: Workbook,
    rng: random.Random,
    sheet: str,
    address: CellAddress | None = None,
    text: str | None = None,
) -> tuple[Workbook, MutationRecord]:
    sh = get_sheet(workbook, sheet)
    if address is None:
        address = CellAddress(sh.n_rows + 2, 0)
    if text is None:
        text = rng.choice(_SCRATCH_TEXTS)

    new_cell = Cell(address=address, value=text, data_type=CellType.STRING)
    new_sheet = set_cell(sh, address, new_cell)
    new_sheet = replace(
        new_sheet, n_rows=max(new_sheet.n_rows, address.row + 1), n_cols=max(new_sheet.n_cols, address.col + 1)
    )
    record = MutationRecord(
        "add_scratch_work", sheet, "excluded_from_grading", {"address": address_to_a1(address), "text": text}
    )
    return replace_sheet(workbook, sheet, new_sheet), record


def _expand_ref(ref: str) -> set[str]:
    """'B3:E3' -> {B3, C3, D3, E3}; 'B3' -> {B3}. Strips any 'Sheet!' prefix."""
    ref = ref.split("!")[-1]
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    assert min_col and min_row and max_col and max_row
    return {
        f"{get_column_letter(c)}{r}"
        for r in range(min_row, max_row + 1)
        for c in range(min_col, max_col + 1)
    }


def wrong_input_correct_method(
    workbook: Workbook,
    rng: random.Random,
    sheet: str,
    input_address: CellAddress | None = None,
    delta: float | None = None,
) -> tuple[Workbook, MutationRecord]:
    """The flagship mutation: break an input value, leave every formula that reads it untouched.

    Uses the Part 2 dependency graph to find which formulas actually depend on
    the chosen input, so the manifest can name them as "correct given inputs" --
    exactly what Part 19's carry-through scoring test needs to check against.
    """
    sh = get_sheet(workbook, sheet)

    reverse: dict[str, list[str]] = {}
    for formula_key, refs in workbook.dependency_graph.items():
        if not formula_key.startswith(f"{sheet}!"):
            continue
        for ref in refs:
            for cell_a1 in _expand_ref(ref):
                reverse.setdefault(cell_a1, []).append(formula_key)

    numeric_inputs = {
        a1: refs
        for a1, refs in reverse.items()
        for addr in [_a1_to_address(a1)]
        if addr in sh.cells and sh.cells[addr].data_type is CellType.NUMBER
    }
    if not numeric_inputs:
        raise ValueError(f"no numeric cell in {sheet!r} is referenced by any formula")

    if input_address is None:
        input_a1 = rng.choice(sorted(numeric_inputs))
        input_address = _a1_to_address(input_a1)
    else:
        input_a1 = address_to_a1(input_address)

    if delta is None:
        delta = rng.choice([-1, 1]) * rng.randint(1, 10)

    cell = sh.cells[input_address]
    if not isinstance(cell.value, int | float):
        raise TypeError(f"{input_address} is not a numeric cell")
    new_value = cell.value + delta
    new_sheet = set_cell(sh, input_address, replace(cell, value=new_value))

    record = MutationRecord(
        "wrong_input_correct_method",
        sheet,
        "carry_through",
        {
            "input_address": input_a1,
            "before": cell.value,
            "after": new_value,
            "formulas_correct_given_bad_input": sorted(numeric_inputs[input_a1]),
        },
    )
    return replace_sheet(workbook, sheet, new_sheet), record


def _a1_to_address(a1: str) -> CellAddress:
    min_col, min_row, _, _ = range_boundaries(a1)
    assert min_col and min_row
    return CellAddress(min_row - 1, min_col - 1)


MUTATIONS: dict[str, Callable[..., tuple[Workbook, MutationRecord]]] = {
    "insert_column": insert_column,
    "delete_row": delete_row,
    "rename_header": rename_header,
    "reorder_columns": reorder_columns,
    "shift_region": shift_region,
    "recolour": recolour,
    "hardcode_formula": hardcode_formula,
    "perturb_value": perturb_value,
    "merge_cells": merge_cells,
    "rename_sheet": rename_sheet,
    "split_table": split_table,
    "add_scratch_work": add_scratch_work,
    "wrong_input_correct_method": wrong_input_correct_method,
}

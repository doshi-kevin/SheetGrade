"""xlsx -> Workbook IR.

Two separate loads of the same file are needed: openpyxl gives you either the
formula string ("=SUM(B3:E3)") or the value Excel last cached for it, never
both from one load. That's the concrete version of "formulas and cached values
can disagree" — openpyxl itself never computes anything, so a formula written
by openpyxl (rather than Excel) has no cached value at all until Excel opens
and recalculates it.
"""

from __future__ import annotations

import datetime
import re
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell as OpenpyxlCell
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from sheetgrade.ir import (
    Borders,
    Cell,
    CellAddress,
    CellType,
    CellValue,
    DefinedName,
    Font,
    Sheet,
    Workbook,
)

# Matches an optional "Sheet1!" prefix followed by a cell or range reference,
# e.g. "B3", "$B$3", "B3:E3", "'My Sheet'!A1". Deliberately simple: this feeds
# the dependency graph (which cells does a formula touch), not a formula
# evaluator, so it doesn't need to understand operators or functions.
_REF_PATTERN = re.compile(
    r"(?:(?:'[^']+')|(?:[A-Za-z_][\w.]*))?!?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?"
)


def read_xlsx(path: Path) -> Workbook:
    formula_wb = openpyxl.load_workbook(path, data_only=False)
    value_wb = openpyxl.load_workbook(path, data_only=True)

    sheets = tuple(
        _read_sheet(formula_ws, value_wb[formula_ws.title]) for formula_ws in formula_wb.worksheets
    )

    defined_names = tuple(
        DefinedName(name=name, refers_to=dn.value) for name, dn in formula_wb.defined_names.items()
    )

    dependency_graph = _build_dependency_graph(sheets)

    return Workbook(sheets=sheets, defined_names=defined_names, dependency_graph=dependency_graph)


def _read_sheet(formula_ws: Worksheet, value_ws: Worksheet) -> Sheet:
    merge_ranges = {str(r) for r in formula_ws.merged_cells.ranges}

    cells: dict[CellAddress, Cell] = {}
    for row in formula_ws.iter_rows():
        for formula_cell in row:
            if isinstance(formula_cell, MergedCell):
                continue  # only the merge's anchor cell (top-left) carries real content
            if formula_cell.value is None and formula_cell.comment is None:
                continue  # skip genuinely empty cells; keeps the IR sparse, not a dense grid

            assert formula_cell.row is not None and formula_cell.column is not None
            value_cell = value_ws.cell(row=formula_cell.row, column=formula_cell.column)
            address = CellAddress(row=formula_cell.row - 1, col=formula_cell.column - 1)
            cells[address] = _read_cell(formula_cell, value_cell.value, address, merge_ranges)

    return Sheet(
        name=formula_ws.title,
        cells=cells,
        n_rows=formula_ws.max_row,
        n_cols=formula_ws.max_column,
    )


def _read_cell(
    formula_cell: OpenpyxlCell,
    cached_raw_value: object,
    address: CellAddress,
    merge_ranges: set[str],
) -> Cell:
    is_formula = formula_cell.data_type == "f"

    data_type = _infer_cell_type(formula_cell, is_formula)
    merge_range = _merge_range_containing(address, merge_ranges)

    return Cell(
        address=address,
        value=None if is_formula else _normalize_value(formula_cell.value, address),
        data_type=data_type,
        formula=_formula_string(formula_cell.value, address) if is_formula else None,
        computed_value=_normalize_value(cached_raw_value, address) if is_formula else None,
        number_format=formula_cell.number_format,
        fill_color_rgb=_solid_fill_rgb(formula_cell),
        font=_read_font(formula_cell),
        borders=_read_borders(formula_cell),
        merge_range=merge_range,
        comment=formula_cell.comment.text if formula_cell.comment else None,
    )


def _infer_cell_type(cell: OpenpyxlCell, is_formula: bool) -> CellType:
    if is_formula:
        return CellType.FORMULA
    if cell.value is None:
        return CellType.BLANK
    if cell.data_type == "b":
        return CellType.BOOLEAN
    if cell.data_type == "e":
        return CellType.ERROR
    if cell.is_date:
        return CellType.DATE
    if cell.data_type == "n":
        return CellType.NUMBER
    return CellType.STRING


def _formula_string(raw: object, address: CellAddress) -> str:
    if isinstance(raw, str):
        return raw
    raise NotImplementedError(
        f"cell {address}: array/data-table formulas are not supported yet (got {type(raw).__name__})"
    )


def _normalize_value(raw: object, address: CellAddress) -> CellValue:
    """Collapse openpyxl's wide value union down to the IR's CellValue type.

    Fails loudly on anything we don't have a mapping for yet (rich text runs,
    array-formula results) rather than silently coercing them to a string --
    per the plan's own rule that adversarial cases should fail loudly.
    """
    if raw is None or isinstance(raw, bool | int | float | str):
        return raw
    if isinstance(raw, Decimal):
        return float(raw)
    if isinstance(raw, datetime.date | datetime.time):
        return raw.isoformat()
    if isinstance(raw, datetime.timedelta):
        return str(raw)
    raise NotImplementedError(f"cell {address}: unsupported value type {type(raw).__name__}")


def _solid_fill_rgb(cell: OpenpyxlCell) -> str | None:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return None
    rgb = fill.fgColor.rgb
    return rgb if isinstance(rgb, str) else None


def _read_font(cell: OpenpyxlCell) -> Font:
    f = cell.font
    color_rgb = f.color.rgb if f.color is not None and isinstance(f.color.rgb, str) else None
    return Font(name=f.name, size=f.size, bold=bool(f.bold), italic=bool(f.italic), color_rgb=color_rgb)


def _read_borders(cell: OpenpyxlCell) -> Borders:
    b = cell.border
    return Borders(
        top=b.top.style if b.top else None,
        bottom=b.bottom.style if b.bottom else None,
        left=b.left.style if b.left else None,
        right=b.right.style if b.right else None,
    )


def _merge_range_containing(address: CellAddress, merge_ranges: set[str]) -> str | None:
    for merge_range in merge_ranges:
        min_col, min_row, max_col, max_row = range_boundaries(merge_range)
        assert min_col is not None and min_row is not None and max_col is not None and max_row is not None
        if min_row - 1 <= address.row <= max_row - 1 and min_col - 1 <= address.col <= max_col - 1:
            return merge_range
    return None


def _build_dependency_graph(sheets: tuple[Sheet, ...]) -> dict[str, frozenset[str]]:
    graph: dict[str, frozenset[str]] = {}
    for sheet in sheets:
        for cell in sheet.cells.values():
            if cell.formula is None:
                continue
            key = f"{sheet.name}!{_address_to_a1(cell.address)}"
            refs = frozenset(_REF_PATTERN.findall(cell.formula))
            graph[key] = refs
    return graph


def _address_to_a1(address: CellAddress) -> str:
    return f"{get_column_letter(address.col + 1)}{address.row + 1}"

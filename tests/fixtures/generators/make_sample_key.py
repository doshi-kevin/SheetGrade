"""Generates tests/fixtures/keys/sample_key.xlsx — a small, hand-crafted instructor
answer-key workbook used as the first real input to the Part 2 xlsx reader.

Re-run with: uv run python tests/fixtures/generators/make_sample_key.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_PATH = Path(__file__).parents[1] / "keys" / "sample_key.xlsx"


def build() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"

    # Merged title row — exercises merge-state handling in the IR.
    ws.merge_cells("A1:E1")
    ws["A1"] = "Quarterly Budget"
    ws["A1"].font = Font(bold=True, size=14)

    headers = ["Category", "Q1", "Q2", "Q3", "Q4", "Total"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=text)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")

    rows = [
        ("Revenue", 1000, 1100, 1200, 1300),
        ("Expenses", 700, 750, 800, 820),
    ]
    for r, (label, *values) in enumerate(rows, start=3):
        ws.cell(row=r, column=1, value=label)
        for c, v in enumerate(values, start=2):
            cell = ws.cell(row=r, column=c, value=v)
            cell.number_format = "$#,##0"
        # Total column: a real formula, not a hardcoded number.
        first_col = get_column_letter(2)
        last_col = get_column_letter(len(values) + 1)
        ws.cell(row=r, column=6, value=f"=SUM({first_col}{r}:{last_col}{r})")

    # Profit row: formula referencing the two rows above it — the kind of
    # dependency the Part 16 scoring graph eventually walks.
    profit_row = 5
    ws.cell(row=profit_row, column=1, value="Profit")
    for c in range(2, 7):
        col_letter = get_column_letter(c)
        cell = ws.cell(row=profit_row, column=c, value=f"={col_letter}3-{col_letter}4")
        cell.font = Font(bold=True)

    ws["A3"].comment = Comment("Confirmed against last year's actuals.", "Instructor")

    return wb


if __name__ == "__main__":
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    build().save(OUTPUT_PATH)
    print(f"wrote {OUTPUT_PATH}")
